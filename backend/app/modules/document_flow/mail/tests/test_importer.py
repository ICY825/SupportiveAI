"""Test đọc file danh sách của lễ tân.

Chạy trên **file mẫu thật** `docs/architecture/mau.xlsx` chứ không chỉ
trên file tự dựng: cả CR-001 tồn tại là vì file thật khác với giả định ban
đầu, nên bài test quan trọng nhất là bài đọc đúng file thật.

Tiêu chí nghiệm thu ở CR-001 §11, phần "Importer".
"""

from __future__ import annotations

import io
from datetime import date, datetime
from pathlib import Path

import pytest

from app.core.exceptions import ValidationError
from app.modules.document_flow.mail import importer

FILE_MAU = Path(__file__).resolve().parents[6] / "docs" / "architecture" / "mau.xlsx"

HEADER = ["stt", "người gửi", "ngày nhận", "số lượng", "nội dung", "người nhận"]


def dung_xlsx(rows, *, header=None, sheet_title="Trang_tính1") -> bytes:
    """Dựng một file `.xlsx` trong bộ nhớ, đúng hình dạng file lễ tân."""
    import openpyxl

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = sheet_title
    sheet.append(header if header is not None else HEADER)
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def dung_csv(rows, *, header=None) -> bytes:
    lines = [",".join(str(c) for c in (header if header is not None else HEADER))]
    lines += [",".join("" if c is None else str(c) for c in row) for row in rows]
    return ("\n".join(lines)).encode("utf-8")


# ----------------------------------------------------------------------
# File mẫu thật
# ----------------------------------------------------------------------

class TestFileMauThat:
    @pytest.fixture
    def report(self):
        if not FILE_MAU.exists():           # pragma: no cover
            pytest.skip(f"Không tìm thấy file mẫu {FILE_MAU}")
        return importer.read_xlsx(FILE_MAU.read_bytes())

    def test_doc_ra_ba_dong(self, report):
        assert len(report.rows) == 3

    def test_so_luong_lan_luot_1_2_1(self, report):
        """CR-001 §11 — con số cụ thể của file mẫu."""
        assert [r.quantity for r in report.rows] == [1, 2, 1]

    def test_o_ngay_datetime_doc_dung(self, report):
        assert all(r.received_on == date(2026, 9, 17) for r in report.rows)
        assert not report.ambiguous_date

    def test_doc_du_sau_cot(self, report):
        dau = report.rows[0]
        assert dau.row_index == 1
        assert dau.sender == "cty nam hải"
        assert dau.content_type == "phong bì"
        assert dau.recipient_name == "Lương hiền"

    def test_ten_nguoi_nhan_giu_nguyen_ban_tho(self, report):
        """Tên viết không chuẩn là dữ liệu, không phải lỗi — đừng "sửa" hộ."""
        assert [r.recipient_name for r in report.rows] == [
            "Lương hiền", "Nguyễn Thị Thu", "Hùng Mạnh"
        ]

    def test_file_khong_co_so_dien_thoai(self, report):
        """Lý do tồn tại của cả CR-001 (§0)."""
        assert all(r.recipient_phone is None for r in report.rows)


# ----------------------------------------------------------------------
# Đọc ngày (CR-001 §3.4)
# ----------------------------------------------------------------------

class TestDocNgay:
    def test_o_datetime_dung_truc_tiep(self):
        parsed = importer.parse_receipt_date(datetime(2026, 9, 17, 14, 30))
        assert parsed.value == date(2026, 9, 17)
        assert not parsed.ambiguous

    def test_chuoi_ro_rang_doc_theo_quy_uoc_viet_nam(self):
        """Ngày 25 không thể là tháng — chỉ một cách đọc hợp lệ."""
        parsed = importer.parse_receipt_date("25/09/2026")
        assert parsed.value == date(2026, 9, 25)
        assert not parsed.ambiguous

    def test_hai_chu_so_nam_hieu_la_the_ky_nay(self):
        assert importer.parse_receipt_date("25/09/26").value == date(2026, 9, 25)

    def test_chuoi_mo_ho_thi_khong_doan(self):
        """`09/17/26` và `17/09/26` lẫn vào nhau — đánh dấu, không đoán bừa.

        Đây là rủi ro thật: định dạng file mẫu đang là `mm-dd-yy`.
        """
        parsed = importer.parse_receipt_date("09/10/2026")
        assert parsed.ambiguous
        # Vẫn ưu tiên quy ước Việt Nam để có cái mà hiện ra cho HC sửa.
        assert parsed.value == date(2026, 10, 9)

    def test_chi_mot_cach_doc_hop_le_thi_khong_mo_ho(self):
        parsed = importer.parse_receipt_date("09/17/2026")
        assert parsed.value == date(2026, 9, 17)
        assert not parsed.ambiguous

    @pytest.mark.parametrize("raw", ["", None, "hôm qua", "17-09", "32/13/2026"])
    def test_ngay_khong_doc_duoc_thi_bao_loi(self, raw):
        with pytest.raises(ValidationError):
            importer.parse_receipt_date(raw)

    def test_lo_bi_danh_dau_khi_co_dong_mo_ho(self):
        content = dung_csv([[1, "cty nam hải", "09/10/2026", 1, "phong bì", "Lương hiền"]])
        report = importer.read_csv(content)
        assert report.ambiguous_date


# ----------------------------------------------------------------------
# Cột và ô thiếu
# ----------------------------------------------------------------------

class TestCotVaOThieu:
    def test_thieu_cot_bat_buoc_thi_bao_ro_cot_nao(self):
        content = dung_xlsx([[1, "cty nam hải", 1, "phong bì"]],
                            header=["stt", "người gửi", "số lượng", "nội dung"])
        with pytest.raises(ValidationError) as exc:
            importer.read_xlsx(content)
        assert "người nhận" in exc.value.message
        assert "ngày nhận" in exc.value.message

    def test_so_luong_de_trong_thi_mac_dinh_1(self):
        """Ô trống nghĩa là lễ tân không buồn ghi "1", không phải không có kiện."""
        content = dung_csv([[1, "cty nam hải", "17/09/2026", "", "phong bì", "Lương hiền"]])
        assert importer.read_csv(content).rows[0].quantity == 1

    def test_so_luong_khong_doc_duoc_thi_bao_loi_kem_so_dong(self):
        content = dung_csv([[1, "cty nam hải", "17/09/2026", "vài cái", "phong bì", "A B"]])
        with pytest.raises(ValidationError, match="Dòng 2"):
            importer.read_csv(content)

    def test_so_luong_am_thi_bao_loi(self):
        content = dung_csv([[1, "cty nam hải", "17/09/2026", -2, "phong bì", "A B"]])
        with pytest.raises(ValidationError):
            importer.read_csv(content)

    def test_dong_trong_thi_bo_qua(self):
        content = dung_xlsx([
            [1, "cty nam hải", datetime(2026, 9, 17), 1, "phong bì", "Lương hiền"],
            [None, None, None, None, None, None],
            [2, "an nguyễn", datetime(2026, 9, 17), 2, "phong bì", "Nguyễn Thị Thu"],
        ])
        assert len(importer.read_xlsx(content).rows) == 2

    def test_file_rong_thi_bao_loi(self):
        with pytest.raises(ValidationError, match="rỗng"):
            importer.read_csv(b"")

    def test_cot_la_thi_bo_qua_khong_bao_loi(self):
        """Lễ tân thêm cột ghi chú thì không được vì thế mà gãy."""
        content = dung_csv(
            [[1, "cty nam hải", "17/09/2026", 1, "phong bì", "Lương hiền", "ghi chú gì đó"]],
            header=HEADER + ["ghi chú"],
        )
        assert importer.read_csv(content).rows[0].recipient_name == "Lương hiền"


class TestTenCot:
    @pytest.mark.parametrize("header", [
        HEADER,
        ["STT", "Người gửi", "Ngày nhận", "Số lượng", "Nội dung", "Người nhận"],
        ["stt", "NGƯỜI GỬI", "ngày nhận:", "số lượng", "nội dung", "người nhận "],
    ])
    def test_hoa_thuong_va_dau_hai_cham_khong_anh_huong(self, header):
        content = dung_csv([[1, "cty nam hải", "17/09/2026", 1, "phong bì", "Lương hiền"]],
                           header=header)
        assert importer.read_csv(content).rows[0].sender == "cty nam hải"

    def test_nhan_san_cot_so_dien_thoai_cho_D7(self):
        """Lễ tân thêm cột thì importer đọc được ngay, không phải sửa code."""
        content = dung_csv(
            [[1, "cty nam hải", "17/09/2026", 1, "phong bì", "Lương hiền", "0912345678"]],
            header=HEADER + ["số điện thoại"],
        )
        assert importer.read_csv(content).rows[0].recipient_phone == "0912345678"


class TestChonBoDoc:
    def test_xlsx(self):
        content = dung_xlsx([[1, "cty nam hải", datetime(2026, 9, 17), 1, "phong bì", "A B"]])
        assert len(importer.read_file("le-tan.xlsx", content).rows) == 1

    def test_csv(self):
        content = dung_csv([[1, "cty nam hải", "17/09/2026", 1, "phong bì", "A B"]])
        assert len(importer.read_file("le-tan.csv", content).rows) == 1

    def test_dinh_dang_la_thi_bao_loi(self):
        with pytest.raises(ValidationError, match="Định dạng không hỗ trợ"):
            importer.read_file("le-tan.pdf", b"...")

    def test_nhieu_sheet_thi_canh_bao(self):
        import openpyxl

        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "Trang_tính1"
        sheet.append(HEADER)
        sheet.append([1, "cty nam hải", datetime(2026, 9, 17), 1, "phong bì", "A B"])
        workbook.create_sheet("Sheet2")
        buffer = io.BytesIO()
        workbook.save(buffer)

        report = importer.read_xlsx(buffer.getvalue())
        assert report.warnings
        assert "sheet" in report.warnings[0]
